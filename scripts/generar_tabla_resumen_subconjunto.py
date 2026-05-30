from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "tabla_resumen_subconjunto_integrado.docx"
FALLBACK_OUT = ROOT / "docs" / "tabla_resumen_subconjunto_integrado_v2.docx"

ROWS = [
    (
        "Dataset completo SQUID-5",
        "138.733 im\u00e1genes TIFF",
        "Referencia general del levantamiento",
    ),
    (
        "Ortomosaico USGS publicado",
        "Aproximadamente 800 \u00d7 160 m; 5 mm",
        "Referencia externa del benchmark",
    ),
    (
        "L\u00edneas revisadas",
        "17\u201324",
        "Delimitaci\u00f3n del universo inicial de selecci\u00f3n",
    ),
    ("Im\u00e1genes iniciales", "4.739", "Universo inicial de selecci\u00f3n"),
    ("Peso aproximado", "35,2 GB", "Estimaci\u00f3n del costo computacional"),
    ("Subbloque seleccionado", "634 im\u00e1genes", "\u00c1rea de trabajo efectiva"),
    ("Im\u00e1genes eliminadas", "38", "Control de calidad preliminar"),
    ("Im\u00e1genes finales", "596", "Procesamiento SfM-MVS y ortomosaico"),
    (
        "Criterio de selecci\u00f3n",
        "Continuidad espacial, solape entre trayectorias y tama\u00f1o manejable",
        "Justificaci\u00f3n metodol\u00f3gica del recorte",
    ),
    (
        "Objetivo del recorte",
        "Reducir costo computacional y evaluar el flujo SfM-MVS con correcci\u00f3n colorim\u00e9trica previa",
        "Prop\u00f3sito experimental del subconjunto",
    ),
]


def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for side, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")

    for border_name in ("top", "bottom"):
        border = borders.find(qn(f"w:{border_name}"))
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def set_cell_bottom_border(cell, size="8", color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for border_name in ("top", "left", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")

    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)


def clear_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def format_run(run, bold=False, italic=False, color="000000"):
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.59)
section.page_height = Cm(27.94)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(11)

caption = doc.add_paragraph()
caption.paragraph_format.space_after = Pt(6)
caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
caption_label = caption.add_run("Tabla X. ")
format_run(caption_label, bold=True)
caption_title = caption.add_run(
    "Resumen integrado de selecci\u00f3n del subconjunto de procesamiento."
)
format_run(caption_title, italic=True)

table = doc.add_table(rows=1, cols=3)
table.autofit = False
set_table_width(table, 9360)
set_table_borders(table)

headers = ("Elemento", "Valor reportado", "Uso en el informe")
widths = (2434, 2995, 3931)

for idx, text in enumerate(headers):
    cell = table.rows[0].cells[idx]
    set_cell_width(cell, widths[idx])
    set_cell_bottom_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    format_run(run, bold=True)

for label, value, use in ROWS:
    cells = table.add_row().cells
    for idx, text in enumerate((label, value, use)):
        cell = cells[idx]
        set_cell_width(cell, widths[idx])
        clear_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        format_run(run)

try:
    doc.save(OUT)
    print(OUT)
except PermissionError:
    doc.save(FALLBACK_OUT)
    print(FALLBACK_OUT)
